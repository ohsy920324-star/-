# -*- coding: utf-8 -*-
"""
강의 정리 store.json -> 완성본 .docx 생성기
예시 파일과 동일한 형식: 표지 -> 목차(필드) -> '이 자료 보는 법' 범례 -> 본문(5색 박스)
"""
import json, sys
from docx import Document
from docx.shared import Pt, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
KFONT = "맑은 고딕"

BOX = {
    "용어":                    {"bg": "E2F0D9", "bd": "70AD47", "label": RGBColor(0x53,0x81,0x35)},
    "사장님께 이렇게 말하세요": {"bg": "FFF2CC", "bd": "BF9000", "label": RGBColor(0x7F,0x60,0x00)},
    "바로 쓰는 실전 멘트":     {"bg": "FCE4EC", "bd": "C5407E", "label": RGBColor(0xAD,0x14,0x57)},
    "핵심":                    {"bg": "DEEBF7", "bd": "2E75B6", "label": RGBColor(0x1F,0x4E,0x79)},
    "주의":                    {"bg": "FCE4D6", "bd": "C55A11", "label": RGBColor(0x83,0x3C,0x00)},
}
LEGEND = [
    ("용어", "마케팅 용어를 쉽게 풀어 설명합니다. 헷갈리는 단어는 여기서 확인하세요."),
    ("사장님께 이렇게 말하세요", "사장님(광고주/대표님)에게 그대로 말로 풀어서 설명할 때 쓰는 멘트 예시입니다."),
    ("바로 쓰는 실전 멘트", "카톡·문의·블로그에 바로 복붙해서 쓸 수 있는 실전 멘트·스크립트, 그리고 사례 글의 원문입니다."),
    ("핵심", "절대 잊으면 안 되는 핵심 포인트입니다."),
    ("주의", "실수하기 쉬운 부분, 꼭 조심해야 할 내용입니다."),
]

def set_kfont(run):
    run.font.name = KFONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'):
        rFonts.set(qn(a), KFONT)

def shade(cell, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),fill)
    cell._tc.get_or_add_tcPr().append(shd)

def cell_borders(cell, color, left_sz=24, other_sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        e = OxmlElement('w:'+edge)
        e.set(qn('w:val'),'single')
        e.set(qn('w:sz'), str(left_sz if edge=='left' else other_sz))
        e.set(qn('w:space'),'0'); e.set(qn('w:color'),color)
        borders.append(e)
    tcPr.append(borders)

def add_box(doc, label, text):
    info = BOX.get(label, BOX["핵심"])
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    cell = tbl.cell(0,0)
    shade(cell, info["bg"]); cell_borders(cell, info["bd"])
    # 라벨
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(4)
    r = p.add_run("[" + label + "]"); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = info["label"]; set_kfont(r)
    # 본문(줄바꿈 분리)
    for i, line in enumerate(str(text).split("\n")):
        cp = cell.add_paragraph(); cp.paragraph_format.space_after = Pt(1)
        rr = cp.add_run(line); rr.font.size = Pt(11); set_kfont(rr)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fldBegin = OxmlElement('w:fldChar'); fldBegin.set(qn('w:fldCharType'),'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve')
    instr.text = r'TOC \o "1-2" \h \z \u'
    fldSep = OxmlElement('w:fldChar'); fldSep.set(qn('w:fldCharType'),'separate')
    t = OxmlElement('w:t'); t.text = "목차를 우클릭 → 필드 업데이트(F9) 하면 페이지 번호가 채워집니다."
    fldEnd = OxmlElement('w:fldChar'); fldEnd.set(qn('w:fldCharType'),'end')
    run._r.append(fldBegin); run._r.append(instr); run._r.append(fldSep); run._r.append(t); run._r.append(fldEnd)

def enable_update_fields(doc):
    settings = doc.settings.element
    el = OxmlElement('w:updateFields'); el.set(qn('w:val'),'true')
    settings.append(el)

def style_doc(doc):
    normal = doc.styles['Normal']
    normal.font.name = KFONT; normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a), KFONT)
    for name, size, color in [('Heading 1',15,NAVY),('Heading 2',13,RGBColor(0x2A,0x3B,0x5E))]:
        st = doc.styles[name]; st.font.name = KFONT; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = color
        rp = st.element.get_or_add_rPr(); rfo = rp.find(qn('w:rFonts'))
        if rfo is None: rfo = OxmlElement('w:rFonts'); rp.append(rfo)
        for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rfo.set(qn(a), KFONT)

def build(store_path, out_path):
    data = json.load(open(store_path, encoding='utf-8'))
    title = data.get("title") or "강의 정리"
    subtitle = data.get("subtitle") or ""
    doc = Document()
    style_doc(doc); enable_update_fields(doc)

    # 표지
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY; set_kfont(r)
    if subtitle:
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(subtitle); r2.font.size = Pt(11); r2.font.color.rgb = RGBColor(0x59,0x59,0x59); set_kfont(r2)
    doc.add_page_break()

    # 목차
    p = doc.add_paragraph(); r = p.add_run("목차"); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY; set_kfont(r)
    add_toc(doc)
    doc.add_page_break()

    # 범례
    doc.add_heading("이 자료 보는 법", level=1)
    for label, desc in LEGEND:
        add_box(doc, label, desc)
    doc.add_page_break()

    # 본문
    for piece in data.get("pieces", []):
        for b in piece.get("blocks", []):
            t = b.get("type"); text = b.get("text","")
            if t == "h1": doc.add_heading(text, level=1)
            elif t == "h2": doc.add_heading(text, level=2)
            elif t == "para":
                pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(6)
                rr = pp.add_run(text); rr.font.size = Pt(11); set_kfont(rr)
            elif t == "bullet":
                pp = doc.add_paragraph(style='List Bullet')
                rr = pp.add_run(text); rr.font.size = Pt(11); set_kfont(rr)
            elif t == "box":
                add_box(doc, b.get("label","핵심"), text)
    doc.save(out_path)
    print("saved:", out_path)

if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
